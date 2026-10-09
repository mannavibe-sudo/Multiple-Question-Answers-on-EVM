"""One-time script: Word file se questions.json banata hai.
Usage: python parse_docx.py "019-EVM_Questions-.docx"
"""
import json, sys
import docx

path = sys.argv[1] if len(sys.argv) > 1 else "019-EVM_Questions-.docx"
d = docx.Document(path)
tbl_q, tbl_a = d.tables[0], d.tables[1] if len(d.tables) > 1 else None

# Answer table: har row mein 4 cells -> (no, ans, no, ans)
answers = {}
if tbl_a is not None:
    for row in tbl_a.rows:
        c = [x.text.strip() for x in row.cells]
        for i in range(0, len(c) - 1, 2):
            if c[i].isdigit() and c[i + 1]:
                answers[int(c[i])] = c[i + 1].upper()

questions = []
for row in tbl_q.rows[1:]:
    num = row.cells[0].text.strip()
    if not num.isdigit():
        continue
    parts = [p.text.strip().rstrip(",") for p in row.cells[1].paragraphs if p.text.strip()]
    q, opts = parts[0], parts[1:]
    n = int(num)
    ans_idx = "ABCD".index(answers[n])
    assert ans_idx < len(opts), (n, opts, answers[n])
    questions.append({"id": n, "question": q, "options": opts, "answer": ans_idx})

json.dump(questions, open("questions.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)
print(len(questions), "questions saved")
